from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = "DoseDeck_Hackathon_Deck.pptx"

BG = RGBColor(6, 24, 29)
BG_2 = RGBColor(10, 34, 39)
CARD = RGBColor(15, 45, 50)
CARD_2 = RGBColor(21, 57, 61)
WHITE = RGBColor(244, 249, 247)
MUTED = RGBColor(161, 184, 181)
MINT = RGBColor(76, 225, 171)
LIME = RGBColor(203, 241, 91)
CORAL = RGBColor(255, 127, 104)
LINE = RGBColor(43, 77, 79)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def rect(slide, x, y, w, h, fill, radius=True, line=None, width=1):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = line or fill
    s.line.width = Pt(width)
    return s


def line(slide, x1, y1, x2, y2, color=LINE, width=1.5):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x1), Inches(y1), Inches(x2-x1), Inches(max(y2-y1, .015)))
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def circle(slide, x, y, d, fill, line_color=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line_color or fill
    return s


def text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
         margin=0, tracking=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run(); run.text = value
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    if tracking is not None:
        run.font._element.set("spc", str(tracking))
    return box


def rich_text(slide, segments, x, y, w, h, size=18, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    for value, color, bold in segments:
        r = p.add_run(); r.text = value
        r.font.name = "Aptos Display"; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    return box


def add_bg(slide, number, section):
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = BG
    circle(slide, 12.45, -0.35, 1.45, BG_2)
    text(slide, "DOSEDECK", .55, .28, 1.9, .25, 10, MINT, True)
    text(slide, section.upper(), 2.35, .28, 3.0, .25, 9, MUTED, True)
    text(slide, f"0{number}", 12.22, 7.05, .55, .2, 9, MUTED, True, align=PP_ALIGN.RIGHT)
    line(slide, .55, 6.98, 11.7, 6.995, LINE, .7)


def pill(slide, label, x, y, w, fill=CARD_2, fg=MINT):
    rect(slide, x, y, w, .34, fill, True)
    text(slide, label.upper(), x, y+.01, w, .22, 9, fg, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def card(slide, x, y, w, h, number, title, body, accent=MINT):
    rect(slide, x, y, w, h, CARD, True, LINE)
    circle(slide, x+.25, y+.25, .43, accent)
    text(slide, number, x+.25, y+.285, .43, .2, 11, BG, True, align=PP_ALIGN.CENTER)
    text(slide, title, x+.25, y+.88, w-.5, .45, 19, WHITE, True)
    text(slide, body, x+.25, y+1.42, w-.5, h-1.65, 12, MUTED)


# Slide 1 — Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = BG
circle(slide, 10.25, -1.1, 4.4, BG_2)
circle(slide, 11.1, -.28, 2.65, CARD)
circle(slide, 11.74, .37, 1.35, MINT)
text(slide, "DD", 11.74, .77, 1.35, .35, 24, BG, True, align=PP_ALIGN.CENTER)
pill(slide, "Hackathon concept", .7, .65, 1.8)
rich_text(slide, [("Dose", WHITE, True), ("Deck", MINT, True)], .7, 1.45, 7.7, 1.0, 52)
text(slide, "The intelligence layer for\neveryday healthcare operations.", .73, 2.5, 8.8, 1.35, 29, WHITE, True)
text(slide, "Turning fragmented signals into trusted, timely action—\nwithout disrupting the workflow people already know.", .73, 4.08, 7.6, .8, 16, MUTED)
line(slide, .73, 5.42, 2.2, 5.445, MINT, 3)
text(slide, "NAVIGATE HEALTHCARE SMARTER", .73, 5.72, 4.4, .3, 11, LIME, True)
text(slide, "CONFIDENTIAL CONCEPT • 2026", .73, 6.75, 3.4, .25, 9, MUTED, True)


# Slide 2 — Problem
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, 2, "The hidden gap")
text(slide, "The problem is not a lack of tools.", .7, .88, 8.8, .6, 31, WHITE, True)
rich_text(slide, [("It is the space ", WHITE, False), ("between them.", MINT, True)], .7, 1.48, 8.8, .55, 28)

card(slide, .7, 2.38, 3.7, 2.55, "01", "Signals arrive fragmented",
     "Products, invoices, stock, sales and expiry cues live in separate moments—and often separate systems.")
card(slide, 4.82, 2.38, 3.7, 2.55, "02", "Decisions arrive late",
     "Teams react after stockouts, waste or reconciliation errors become visible.", LIME)
card(slide, 8.94, 2.38, 3.7, 2.55, "03", "Trust remains manual",
     "Automation without verification creates risk; manual checks create friction.", CORAL)
rect(slide, .7, 5.35, 11.94, .95, CARD_2, True, LINE)
text(slide, "The opportunity", .98, 5.63, 1.6, .25, 11, MINT, True)
text(slide, "Create one trusted decision layer across the operational flow.", 2.55, 5.54, 9.5, .4, 19, WHITE, True)


# Slide 3 — Concept
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, 3, "The concept")
text(slide, "One calm surface. Three intelligent moves.", .7, .88, 10.4, .55, 31, WHITE, True)
text(slide, "DoseDeck observes the workflow, understands context, and helps the operator act—with confirmation built in.", .7, 1.52, 11.4, .55, 15, MUTED)

xs = [.7, 4.63, 8.56]
titles = ["OBSERVE", "UNDERSTAND", "ORCHESTRATE"]
subtitles = ["Capture operational signals", "Reveal what matters now", "Guide the next best action"]
symbols = ["◉", "◇", "→"]
colors = [MINT, LIME, CORAL]
for i, x in enumerate(xs):
    rect(slide, x, 2.42, 3.45, 2.62, CARD, True, LINE)
    text(slide, symbols[i], x+.28, 2.67, .55, .55, 26, colors[i], True)
    text(slide, titles[i], x+.28, 3.36, 2.8, .3, 13, colors[i], True)
    text(slide, subtitles[i], x+.28, 3.83, 2.75, .72, 19, WHITE, True)
    if i < 2:
        text(slide, "+", x+3.59, 3.48, .28, .35, 22, MUTED, True, align=PP_ALIGN.CENTER)

rect(slide, 2.16, 5.48, 9.02, .72, BG_2, True, LINE)
for i, label in enumerate(["Fewer blind spots", "Lower avoidable waste", "Faster confident decisions"]):
    px = 2.45 + i*2.92
    circle(slide, px, 5.74, .12, MINT)
    text(slide, label, px+.22, 5.62, 2.55, .25, 11, WHITE, True)


# Slide 4 — Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, 4, "Architecture")
text(slide, "Designed as a trusted intelligence layer.", .7, .82, 10.7, .55, 30, WHITE, True)
text(slide, "A modular, model-agnostic architecture keeps the experience simple while the system learns behind the scenes.", .7, 1.42, 11.5, .45, 14, MUTED)

# left input rail
text(slide, "SIGNALS", .73, 2.19, 1.25, .25, 10, MINT, True)
for i, label in enumerate(["Visual", "Inventory", "Transactions", "Operator"]):
    y = 2.63 + i*.72
    rect(slide, .7, y, 1.62, .48, CARD, True, LINE)
    text(slide, label, .7, y+.12, 1.62, .2, 11, WHITE, True, align=PP_ALIGN.CENTER)

# center stack
text(slide, "INTELLIGENCE CORE", 3.0, 2.19, 2.4, .25, 10, LIME, True)
layers = [
    ("Context engine", "Interprets what is happening", CARD_2),
    ("Decision engine", "Ranks risk and opportunity", CARD),
    ("Workflow engine", "Turns insight into guided action", CARD_2),
]
for i, (t, b, fc) in enumerate(layers):
    y = 2.62 + i*1.04
    rect(slide, 2.98, y, 4.15, .78, fc, True, LINE)
    text(slide, t, 3.22, y+.13, 1.95, .25, 15, WHITE, True)
    text(slide, b, 5.1, y+.16, 1.8, .28, 10, MUTED, False, align=PP_ALIGN.RIGHT)

# output rail
text(slide, "ACTIONS", 7.84, 2.19, 1.25, .25, 10, CORAL, True)
for i, label in enumerate(["Alert", "Recommend", "Confirm", "Audit"]):
    y = 2.63 + i*.72
    rect(slide, 7.82, y, 1.62, .48, CARD, True, LINE)
    text(slide, label, 7.82, y+.12, 1.62, .2, 11, WHITE, True, align=PP_ALIGN.CENTER)

# trust rail
rect(slide, 10.17, 2.16, 2.47, 3.36, BG_2, True, LINE)
text(slide, "TRUST LAYER", 10.46, 2.49, 1.75, .25, 10, MINT, True)
for i, (a,b) in enumerate([
    ("Human-in-loop", "before change"),
    ("Explainability", "behind advice"),
    ("Auditability", "after action"),
    ("Local control", "over operations"),
]):
    y = 2.93 + i*.59
    circle(slide, 10.47, y+.03, .13, MINT)
    text(slide, a, 10.73, y, 1.52, .2, 11, WHITE, True)
    text(slide, b, 10.73, y+.21, 1.52, .18, 9, MUTED)

# arrows + foundation
text(slide, "→", 2.42, 3.52, .42, .4, 22, MUTED, True, align=PP_ALIGN.CENTER)
text(slide, "→", 7.27, 3.52, .42, .4, 22, MUTED, True, align=PP_ALIGN.CENTER)
rect(slide, .7, 5.83, 11.94, .55, CARD_2, True, LINE)
text(slide, "SECURE OPERATIONAL DATA FABRIC  •  API-FIRST  •  MODEL-AGNOSTIC  •  MODULAR", .7, 5.99, 11.94, .2, 10, LIME, True, align=PP_ALIGN.CENTER)


# Slide 5 — Moat
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, 5, "Defensibility")
text(slide, "Built to compound. Difficult to displace.", .7, .84, 10.9, .55, 31, WHITE, True)
text(slide, "Scale can copy features. It cannot instantly copy context, trust, or embedded workflow memory.", .7, 1.46, 11.3, .45, 15, MUTED)

# flywheel
cx, cy = 3.18, 4.14
circle(slide, 2.2, 3.16, 1.96, CARD_2)
text(slide, "DOSEDECK\nCONTEXT\nGRAPH", 2.2, 3.63, 1.96, .72, 15, MINT, True, align=PP_ALIGN.CENTER)
nodes = [
    (2.44, 2.05, "Verified\nactions"),
    (4.25, 3.54, "Better\ndecisions"),
    (2.46, 5.19, "Deeper\nworkflow fit"),
    (.62, 3.54, "Richer local\ncontext"),
]
for x,y,label in nodes:
    circle(slide, x, y, 1.42, CARD, MINT)
    text(slide, label, x+.08, y+.42, 1.26, .52, 11, WHITE, True, align=PP_ALIGN.CENTER)
text(slide, "↻", 2.61, 3.72, 1.15, .65, 30, LIME, True, align=PP_ALIGN.CENTER)

# pillars
text(slide, "WHY A LARGE PLATFORM DOESN’T WIN BY DEFAULT", 6.12, 2.12, 5.7, .3, 11, CORAL, True)
pillars = [
    ("01", "Local context graph", "Each verified action sharpens site-specific understanding."),
    ("02", "Workflow memory", "Value accumulates inside the operator’s daily rhythm."),
    ("03", "Trust infrastructure", "Explanations, confirmations and audit trails earn adoption."),
    ("04", "Switching gravity", "History, preferences and integrations deepen over time."),
]
for i, (n,t,b) in enumerate(pillars):
    y = 2.65 + i*.87
    rect(slide, 6.1, y, 6.12, .68, CARD if i%2==0 else BG_2, True, LINE)
    text(slide, n, 6.35, y+.18, .4, .2, 10, MINT, True)
    text(slide, t, 6.88, y+.13, 1.75, .22, 13, WHITE, True)
    text(slide, b, 8.7, y+.12, 3.2, .35, 10, MUTED)

text(slide, "Defensibility thesis—not an absolute claim. The moat grows with verified use and integration depth.", 6.12, 6.22, 6.0, .3, 9, MUTED)


# Slide 6 — Close
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, 6, "The opportunity")
pill(slide, "The next operating layer", .7, .84, 2.2)
text(slide, "From reactive operations\nto a learning system.", .7, 1.47, 8.2, 1.22, 35, WHITE, True)
text(slide, "DoseDeck starts with the moments where clarity matters most—then compounds into an operational advantage that gets stronger with every verified action.", .7, 2.95, 7.25, 1.0, 16, MUTED)

metrics = [
    ("SEE", "Blind spots earlier"),
    ("DECIDE", "With confidence"),
    ("ACT", "Without disruption"),
]
for i, (a,b) in enumerate(metrics):
    x = .7 + i*2.55
    rect(slide, x, 4.43, 2.25, 1.14, CARD, True, LINE)
    text(slide, a, x+.22, 4.68, 1.8, .22, 11, MINT if i<2 else LIME, True)
    text(slide, b, x+.22, 5.02, 1.78, .25, 12, WHITE, True)

rect(slide, 9.05, 1.29, 3.42, 4.58, CARD_2, True, LINE)
circle(slide, 10.04, 1.87, 1.44, MINT)
text(slide, "DD", 10.04, 2.28, 1.44, .34, 24, BG, True, align=PP_ALIGN.CENTER)
text(slide, "DoseDeck", 9.45, 3.65, 2.62, .4, 24, WHITE, True, align=PP_ALIGN.CENTER)
text(slide, "NAVIGATE\nHEALTHCARE\nSMARTER", 9.45, 4.28, 2.62, .9, 13, LIME, True, align=PP_ALIGN.CENTER)
text(slide, "Let’s build the intelligence\nbehind better operations.", 8.97, 6.25, 3.65, .48, 13, WHITE, True, align=PP_ALIGN.CENTER)


# Core metadata
prs.core_properties.title = "DoseDeck — Hackathon Concept"
prs.core_properties.subject = "Six-slide concept, architecture and defensibility deck"
prs.core_properties.author = "DoseDeck Team"
prs.core_properties.keywords = "DoseDeck, healthcare operations, AI, architecture, hackathon"
prs.core_properties.comments = "High-level concept deck. Implementation details intentionally abstracted."

prs.save(OUT)
print(OUT)
